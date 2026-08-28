from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import func
from core.permissions import require_admin_access, require_write_access
from core.database import get_db
from core.security import get_current_user
from models.entry import DocumentEntry
from models.history import DocumentHistory
from models.register_type import RegisterType
from models.user import User
from schemas.entry import EntryCreate, EntryResponse, StatusUpdate, PendingNotificationResponse
from typing import List
import openpyxl
from openpyxl.styles import Border, Side, Font, Alignment
from io import BytesIO
from models.file import UploadedFile
from datetime import date
from botocore.exceptions import ClientError
from services.s3 import delete_file_from_s3
router = APIRouter()

LEGACY_FIELD_CONFIG = [
    {"key": "nr_inregistrare", "label": "Nr Inregistrare", "source": "column"},
    {"key": "cod", "label": "Cod", "source": "column"},
    {"key": "iesire_intrare", "label": "Iesire/Intrare", "source": "column"},
    {"key": "nr_data_emitent", "label": "Nr/Data Emitent", "source": "column"},
    {"key": "nr_data_destinator", "label": "Nr/Data Destinator", "source": "column"},
    {"key": "rezumat", "label": "Rezumat", "source": "column"},
    {"key": "repartizare_dpt", "label": "Repartizare Dpt", "source": "column"},
    {"key": "mod_rezolvare", "label": "Mod Rezolvare", "source": "column"},
    {"key": "nr_pag_anexe", "label": "Nr Pag & Nr Anexe", "source": "column"},
]
TITLE = [
    "Nume Firma",
    "Nume Registru"
]

COLUMNS = [
    "Nr Inregistrare", "Cod", "Iesire/Intrare", "Nr/Data Emitent",
    "Nr/Data Destinator", "Rezumat", "Repartizare Dpt",
    "Mod Rezolvare", "Nr Pag & Nr Anexe"
]

FIELD_MAP = [
    "nr_inregistrare", "cod", "iesire_intrare", "nr_data_emitent",
    "nr_data_destinator", "rezumat", "repartizare_dpt",
    "mod_rezolvare", "nr_pag_anexe"
]

@router.get("/pending", response_model=List[PendingNotificationResponse])
def get_pending_validations(
    db: Session = Depends(get_db), 
    current_user: dict = Depends(get_current_user)
):
    require_admin_access(current_user)
    
    # Perform a join to grab the document AND the user who uploaded it
    pending_docs = db.query(DocumentEntry, User).join(
        User, DocumentEntry.user_id == User.id
    ).filter(
        DocumentEntry.status == "pending"
    ).all()
    
    # Format it for the frontend
    results = []
    for doc, user in pending_docs:
        results.append({
            "id": doc.id,
            "entry_id": doc.id,
            "username": user.username,
            "register_type": doc.register_type,
            "created_at": doc.created_at,
            "nr_inregistrare": doc.nr_inregistrare
        })
        
    return results

@router.post("/", response_model=EntryResponse)
def create_entry(
    entry: EntryCreate,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    require_write_access(current_user)
    
    initial_status = "validated" if current_user.get("permission_level") == "full_access" else "pending"
    
    db_entry = DocumentEntry(**entry.model_dump(), user_id=current_user["id"], status=initial_status)
    db.add(db_entry)
    db.flush()  # assigns db_entry.id without ending the transaction

    # Compatibilitate cu fluxul vechi: dacă a venit un file_id la creare
    # (upload înainte de entry), leagă fișierul respectiv și prin noua relație entry_id,
    # ca să apară și el în GET /files/by-entry/{id}, nu doar în entry.file_id.
    if db_entry.file_id is not None:
        linked_file = db.query(UploadedFile).filter(UploadedFile.id == db_entry.file_id).first()
        if linked_file is not None:
            linked_file.entry_id = db_entry.id

    history_event = DocumentHistory(
        entry_id=db_entry.id,
        user_id=current_user["id"],
        status=initial_status,
    )
    db.add(history_event)

    db.commit()
    db.refresh(db_entry)
    return db_entry
@router.delete("/{entry_id}")
def delete_entry(
    entry_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    require_write_access(current_user)

    db_entry = (
        db.query(DocumentEntry)
        .join(User, User.id == DocumentEntry.user_id)
        .filter(DocumentEntry.id == entry_id, User.company_id == current_user["company_id"])
        .first()
    )
    if not db_entry:
        raise HTTPException(status_code=404, detail="Entry not found or unauthorized")

    attached_files = db.query(UploadedFile).filter(UploadedFile.entry_id == entry_id).all()

    # Detaseaza legatura legacy inainte de a sterge fisierele, altfel FK-ul
    # document_entries_file_id_fkey blocheaza DELETE-ul pe files de mai jos.
    db_entry.file_id = None
    db.flush()

    for file_record in attached_files:
        try:
            delete_file_from_s3(file_record.s3_key)
        except ClientError as e:
            error_code = e.response.get("Error", {}).get("Code") if hasattr(e, "response") else None
            if error_code not in {"NoSuchKey", "404"}:
                raise HTTPException(status_code=500, detail=f"Could not delete file {str(e)}")
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Could not delete file {str(e)}")
        db.delete(file_record)

    db.delete(db_entry)
    db.commit()

    return {"status": "success", "message": "Entry and its files were deleted successfully"}
@router.patch("/{entry_id}/status", response_model=EntryResponse)
def update_entry_status(
    entry_id: int,
    payload: StatusUpdate,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    # Enforce that ONLY admins can change statuses
    from core.permissions import require_admin_access
    require_admin_access(current_user)

    db_entry = db.query(DocumentEntry).filter(DocumentEntry.id == entry_id).first()
    if not db_entry:
        raise HTTPException(status_code=404, detail="Entry not found")

    # Prevent unnecessary database hits if status isn't changing
    if db_entry.status == payload.status:
        return db_entry

    # Update the document's state
    db_entry.status = payload.status

    # Add the validation action to the audit history
    history_event = DocumentHistory(
        entry_id=db_entry.id,
        user_id=current_user["id"],
        status=payload.status,
    )
    db.add(history_event)

    db.commit()
    db.refresh(db_entry)
    return db_entry


@router.get("/next-nr/{register_type:path}")
def get_next_nr(
    register_type: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    today = date.today()
    date_prefix = today.strftime("%Y/%m/%d")
    date_compact = today.strftime("%Y%m%d")

    # count how many entries exist for this user + register + today
    count_today = db.query(DocumentEntry).filter(
        DocumentEntry.user_id == current_user["id"],
        DocumentEntry.register_type == register_type,
        func.date(DocumentEntry.created_at) == today
    ).count()

    next_count = count_today + 1
    count_str = str(next_count).zfill(3)  # 001, 002, etc.

    return {
        "nr_inregistrare": f"{date_prefix}/{count_str}",
        "cod": f"{date_compact}{count_str}"
    }

def sanitize_sheet_title(value: str) -> str:
    invalid_chars = r'\/*?:[]'
    cleaned = "".join(c for c in value if c not in invalid_chars)
    return cleaned[:31] or "Sheet1"


def _get_entry_value(entry: DocumentEntry, field: dict):
    if field["source"] == "custom":
        return (entry.custom_fields or {}).get(field["key"], "")
    return getattr(entry, field["key"], "")

@router.get("/{register_type:path}/excel")
def download_excel(
    register_type: str,
    start_date: date | None = None,
    end_date: date | None = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    if start_date and end_date and start_date > end_date:
        raise HTTPException(status_code=400, detail="Start date cannot be after end date")

    query = (
        db.query(DocumentEntry)
        .join(User, User.id == DocumentEntry.user_id)
        .filter(
            User.company_id == current_user["company_id"],
            DocumentEntry.register_type == register_type,
            DocumentEntry.status != "rejected",
        )
    )

    if start_date is not None:
        query = query.filter(func.date(DocumentEntry.created_at) >= start_date)
    if end_date is not None:
        query = query.filter(func.date(DocumentEntry.created_at) <= end_date)

    entries = query.all()

    register_config = (
        db.query(RegisterType)
        .filter(
            RegisterType.company_id == current_user["company_id"],
            RegisterType.key == register_type,
        )
        .first()
    )
    field_config = register_config.field_config if register_config else LEGACY_FIELD_CONFIG
    display_name = register_config.display_name if register_config else register_type

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_title(register_type)
    num_cols = len(field_config)
    mid = num_cols // 2

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=mid)
    ws.merge_cells(start_row=1, start_column=mid + 1, end_row=1, end_column=num_cols)

    company_cell = ws.cell(row=1, column=1, value=current_user["company_name"])
    company_cell.font = Font(size=16, bold=True)
    company_cell.alignment = Alignment(horizontal="center", vertical="center")

    register_cell = ws.cell(row=1, column=mid + 1, value=display_name)
    register_cell.font = Font(size=16, bold=True)
    register_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 8

    #row 3 - headers
    for col_idx, field in enumerate(field_config, start=1):
        cell = ws.cell(row=3, column=col_idx, value=field["label"])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    #row 4 and below, data
    for row_offset, entry in enumerate(entries, start=4):
        for col_idx, field in enumerate(field_config, start=1):
            cell = ws.cell(row=row_offset, column=col_idx, value=_get_entry_value(entry, field))
            cell.border = border

    # auto-width columns, based on headers + data only (skip title row)
    for col_idx in range(1, num_cols + 1):
        col_letter = ws.cell(row=3, column=col_idx).column_letter
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(3, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[col_letter].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    filename = f"{register_type.replace('/', '-')}.xlsx"
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/{register_type:path}", response_model=List[EntryResponse])
def get_entries(
    register_type: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    return (
        db.query(DocumentEntry)
        .join(User, User.id == DocumentEntry.user_id)
        .filter(
            User.company_id == current_user["company_id"],
            DocumentEntry.register_type == register_type,
        )
        .all()
    )
